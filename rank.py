#!/usr/bin/env python3
"""
rank.py — ARIA v2: Adaptive Retrieval Intelligence Architecture
Redrob Hackathon 2026 | Track 1 — Intelligent Candidate Ranking

5 Novel Techniques:
  1. Career DNA Fingerprinting   — trajectory sequence matching
  2. Implicit Skill Graph        — cluster depth, not keyword presence
  3. Velocity-Weighted Signals   — time-decay on all 23 behavioral signals
  4. Peer Benchmark Percentile   — relative rank within cohort
  5. JD Gap Vector               — requirement-level gap analysis

Usage:
    python rank.py --candidates ./candidates.jsonl --out ./submission.csv
    python rank.py --candidates ./sample_candidates.json --out ./submission.csv
"""

import argparse
import csv
import os
import sys
import time
from functools import cmp_to_key

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from scorer.dna import dna_similarity_score
from scorer.skill_graph import skill_graph_score
from scorer.velocity import velocity_score
from scorer.peer_benchmark import compute_peer_percentiles, peer_context
from scorer.gap_vector import compute_gap_vector
from utils.loader import load_all_candidates, find_candidates_file
from utils.reasoning import generate_reasoning


# ── Final score weights ────────────────────────────────────────────────────
# Tuned to match JD's stated priorities
WEIGHTS = {
    "gap_vector":       0.28,   # Most comprehensive: requirement-level analysis
    "skill_graph":      0.22,   # Cluster depth — are they deep in retrieval?
    "dna":              0.18,   # Career trajectory sequence
    "peer_percentile":  0.16,   # How good are they within their cohort?
    "velocity":         0.16,   # Are they actually available right now?
}


def score_one(candidate: dict) -> dict:
    """Compute all 5 technique scores for one candidate."""

    # ── Hard disqualifier: non-technical career ────────────────────────────
    profile = candidate.get("profile", {})
    history = candidate.get("career_history", [])
    current_title = profile.get("current_title", "").lower()
    hard_non_tech = [
        "project manager", "marketing manager", "sales executive",
        "customer support", "brand design", "graphic design",
        "mechanical engineer", "civil engineer", "accountant",
        "hr manager", "operations manager", "content writer",
    ]
    all_titles = [j.get("title", "").lower() for j in history] + [current_title]
    non_tech_count = sum(
        1 for t in all_titles if any(nt in t for nt in hard_non_tech)
    )
    if len(all_titles) > 0 and non_tech_count / len(all_titles) >= 0.6:
        return {
            "final_score": 0.001,
            "disqualified": True,
            "reason": "Non-technical career profile",
            "dna": {}, "skill_graph": {}, "velocity": {},
            "gap_vector": {}, "peer_context": {},
        }

    # ── Run all 5 scorers ──────────────────────────────────────────────────
    dna     = dna_similarity_score(candidate)
    graph   = skill_graph_score(candidate)
    vel     = velocity_score(candidate)
    gap     = compute_gap_vector(candidate)
    pc      = peer_context(candidate)

    # Honeypot detection via gap vector (no production evidence + implausible skills)
    if gap.get("n_hard_gaps", 0) >= 3:
        # Missing all core requirements — likely honeypot or irrelevant
        # Don't zero out — just let score be very low naturally
        pass

    return {
        "final_score": 0.0,   # filled in after peer percentile pass
        "disqualified": False,
        "dna":          dna,
        "skill_graph":  graph,
        "velocity":     vel,
        "gap_vector":   gap,
        "peer_context": pc,
    }


def rank_candidates(candidates: list, verbose: bool = True) -> list:
    """Full ARIA v2 pipeline. Returns top-100 list sorted by score."""
    t0 = time.time()

    if verbose:
        print(f"\n{'='*60}")
        print(f"  ARIA v2 — Adaptive Retrieval Intelligence Architecture")
        print(f"  Candidates: {len(candidates):,}")
        print(f"{'='*60}")

    # ── Pass 1: Score all candidates with techniques 1,2,3,5 ──────────────
    if verbose:
        print("Pass 1/3: Scoring DNA + Graph + Velocity + Gap Vector...")

    scored = []
    for i, c in enumerate(candidates):
        s = score_one(c)
        scored.append({"candidate": c, "scores": s})
        if verbose and (i + 1) % 20000 == 0:
            print(f"  {i+1:,}/{len(candidates):,}...")

    if verbose:
        print(f"  Done in {time.time()-t0:.1f}s")

    # ── Pass 2: Peer Benchmark Percentile (needs all scores first) ─────────
    if verbose:
        print("Pass 2/3: Computing peer benchmark percentiles...")

    # Compute a raw pre-score for peer grouping
    raw_scores = []
    for item in scored:
        s = item["scores"]
        if s.get("disqualified"):
            raw_scores.append(0.0)
        else:
            raw = (
                WEIGHTS["gap_vector"]  * s["gap_vector"].get("gap_score", 0) +
                WEIGHTS["skill_graph"] * s["skill_graph"].get("score", 0) +
                WEIGHTS["dna"]         * s["dna"].get("score", 0) +
                WEIGHTS["velocity"]    * s["velocity"].get("velocity_score", 0)
            )
            raw_scores.append(raw)

    all_candidates = [item["candidate"] for item in scored]
    peer_percentiles = compute_peer_percentiles(all_candidates, raw_scores)

    # ── Pass 3: Compute final scores ───────────────────────────────────────
    if verbose:
        print("Pass 3/3: Computing final scores + reasoning...")

    for i, item in enumerate(scored):
        s = item["scores"]
        if s.get("disqualified"):
            item["final_score"] = 0.001
            continue

        peer_pct = peer_percentiles[i]
        raw = raw_scores[i]

        final = (
            raw +
            WEIGHTS["peer_percentile"] * peer_pct
        )

        # Velocity as a multiplier (unavailable candidate = deprioritize)
        vel_score = s["velocity"].get("velocity_score", 0.5)
        # Soft multiplier: 0.7 + 0.3 * velocity (never fully zeroes out a good candidate)
        vel_multiplier = 0.70 + 0.30 * vel_score

        final = max(0.0, min(1.0, final * vel_multiplier))
        item["final_score"] = final
        s["peer_percentile"] = peer_pct
        s["final_score"] = final

    # ── Sort + select top 100 ──────────────────────────────────────────────
    def compare(a, b):
        if a["final_score"] != b["final_score"]:
            return -1 if a["final_score"] > b["final_score"] else 1
        aid = a["candidate"]["candidate_id"]
        bid = b["candidate"]["candidate_id"]
        return -1 if aid < bid else 1

    scored.sort(key=cmp_to_key(compare))
    top_100 = scored[:100]

    # ── Generate reasoning + build output ─────────────────────────────────
    results = []
    for rank_idx, item in enumerate(top_100, start=1):
        candidate = item["candidate"]
        s = item["scores"]
        reasoning = generate_reasoning(candidate, s, rank_idx)
        results.append({
            "candidate_id": candidate["candidate_id"],
            "rank": rank_idx,
            "score": item["final_score"],
            "reasoning": reasoning,
        })

    # Enforce monotone non-increasing scores
    for i in range(1, len(results)):
        if results[i]["score"] > results[i-1]["score"]:
            results[i]["score"] = results[i-1]["score"]

    if verbose:
        total = time.time() - t0
        print(f"\n{'='*60}")
        print(f"  Done in {total:.1f}s | Top 3:")
        for r in results[:3]:
            print(f"  #{r['rank']} {r['candidate_id']} | {r['score']:.4f}")
            print(f"     {r['reasoning'][:100]}...")
        print(f"{'='*60}\n")

    return results


def write_submission(results: list, output_path: str):
    """Write results to XLSX (portal requires) or CSV as fallback."""
    import pandas as pd
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

    rows = []
    for r in results:
        rows.append({
            "candidate_id": r["candidate_id"],
            "rank":         r["rank"],
            "score":        round(r["score"], 6),
            "reasoning":    str(r["reasoning"]).replace("\n", " ").replace("\r", " "),
        })

    if output_path.endswith(".xlsx"):
        df = pd.DataFrame(rows)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer_xl:
            df.to_excel(writer_xl, index=False, sheet_name="Rankings")
            ws = writer_xl.sheets["Rankings"]

            # Header styling
            header_fill  = PatternFill("solid", fgColor="4B0082")
            header_font  = Font(bold=True, color="FFFFFF", size=11)
            center_align = Alignment(horizontal="center", vertical="center")
            left_align   = Alignment(horizontal="left", vertical="center", wrap_text=True)
            thin          = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"),  bottom=Side(style="thin"),
            )
            for cell in ws[1]:
                cell.fill = header_fill; cell.font = header_font
                cell.alignment = center_align; cell.border = thin

            # Column widths
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 8
            ws.column_dimensions["C"].width = 12
            ws.column_dimensions["D"].width = 80

            # Alternating row colours
            fill_odd  = PatternFill("solid", fgColor="F3EFFF")
            fill_even = PatternFill("solid", fgColor="FFFFFF")
            for i, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), 1):
                fill = fill_odd if i % 2 else fill_even
                for cell in row:
                    cell.fill = fill; cell.border = thin
                    cell.alignment = left_align if cell.column == 4 else center_align

            ws.freeze_panes = "A2"

            # About sheet
            ws2 = writer_xl.book.create_sheet("About ARIA")
            ws2["A1"] = "ARIA v2 — Adaptive Retrieval Intelligence Architecture"
            ws2["A1"].font = Font(bold=True, size=14, color="4B0082")
            ws2["A2"] = "Redrob Hackathon 2026 | Modern College of Engineering, Pune (SPPU)"
            ws2["A4"] = "5 Novel Techniques:"
            ws2["A4"].font = Font(bold=True)
            for j, t in enumerate([
                "1. Career DNA Fingerprinting  — trajectory sequence (R=Retrieval M=ML P=Platform S=Software D=Data C=Services)",
                "2. Implicit Skill Graph        — cluster depth, not keyword presence",
                "3. Velocity-Weighted Signals   — exponential time-decay on all 23 behavioral signals",
                "4. Peer Benchmark Percentile   — relative rank within YOE x company-tier cohort",
                "5. JD Gap Vector               — requirement-level gap analysis with evidence per candidate",
            ], 5):
                ws2[f"A{j}"] = t
            ws2.column_dimensions["A"].width = 95

        print(f"✅ Submission written to: {output_path} (formatted XLSX)")
    else:
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
            writer.writerow(["candidate_id", "rank", "score", "reasoning"])
            for r in rows:
                writer.writerow([r["candidate_id"], r["rank"], r["score"], r["reasoning"]])
        print(f"✅ Submission written to: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="ARIA v2 — Redrob Candidate Ranker")
    parser.add_argument("--candidates", "-c", type=str, default=None)
    parser.add_argument("--out", "-o", type=str, default="submission.xlsx")
    parser.add_argument("--data-dir", "-d", type=str, default=".")
    parser.add_argument("--limit", "-n", type=int, default=None,
                        help="Limit candidates for testing (e.g. --limit 500)")
    parser.add_argument("--quiet", "-q", action="store_true")
    args = parser.parse_args()

    candidates_path = args.candidates or find_candidates_file(args.data_dir)
    print(f"Loading: {candidates_path}")

    candidates = load_all_candidates(candidates_path, max_count=args.limit,
                                     verbose=not args.quiet)
    results = rank_candidates(candidates, verbose=not args.quiet)
    write_submission(results, args.out)

    # Validate
    try:
        from validate_submission import validate_submission
        errors = validate_submission(args.out)
        if errors:
            print(f"⚠️  Validation issues ({len(errors)}):")
            for e in errors[:5]:
                print(f"   - {e}")
        else:
            print("✅ Submission validated successfully.")
    except Exception:
        print("(validator not found — skipping)")


if __name__ == "__main__":
    main()
